"""
Lê a planilha "Seminovos Tracker" (preço anunciado dos seminovos da Localiza e FIPE, mês a mês) e
gera o JSON que o painel mostra nas abas FIPE e Radar de seminovos.

A planilha é material de terceiros: o JSON vai só para o banco do painel (debentures/tracker_seminovos),
nunca para o repositório.

Blocos lidos:
- aba "Cover - Listing Delta": variação mensal do preço anunciado (consolidado, por tipo de carro,
  por ano-modelo e por modelo);
- aba "Cover - FIPE": spread preço anunciado × FIPE no mesmo mês (T_0), por tipo e por ano-modelo,
  e a variação média mensal da FIPE (consolidado e por tipo).

Os modelos são agrupados por carroceria (Hatch, Sedã, SUV, Picape, Utilitário), com a média simples
das variações do mês; variações acima de 10% em módulo são descartadas da média (anúncio isolado).

Uso: python tools/tracker_seminovos.py <planilha.xlsx> <saida.json> [--base <json anterior>]
     Com --base, os meses que já estavam no JSON anterior e não vêm na planilha nova ficam (histórico mensal).
"""

import json
import re
import statistics
import sys
from datetime import datetime

import openpyxl

MESES = {"jan": 1, "feb": 2, "fev": 2, "mar": 3, "apr": 4, "abr": 4, "may": 5, "mai": 5, "jun": 6, "jul": 7,
         "aug": 8, "ago": 8, "sep": 9, "set": 9, "oct": 10, "out": 10, "nov": 11, "dec": 12, "dez": 12}
TIPOS = {"1Low Entry": "Entrada", "2Mid-size": "Médio", "3Luxury": "Luxo", "4Light Commercial": "Comercial leve"}
CARROCERIA = {
    "Hatch": ["ONIX", "C3", "208", "HB20", "ARGO", "POLO", "MOBI", "KWID", "GOL", "A3"],
    "Sedã": ["VOYAGE", "HB20S", "CRONOS", "VIRTUS", "CITY", "SENTRA", "COROLLA", "CIVIC", "320I", "A4", "530E"],
    "SUV": ["C3 AIRCROSS", "SPORTAGE", "TIGGO 8", "TAOS", "ZR-V", "RENEGADE", "TRACKER", "COMPASS", "TIGUAN", "HR-V",
            "T-CROSS", "PULSE", "DUSTER", "CRETA", "KARDIAN", "FASTBACK", "NIVUS", "KICKS", "COROLLA CROSS", "2008",
            "COMMANDER", "XC 60", "X1", "Q5", "HAVAL", "EQUINOX", "E-TRON", "TRAILBLAZER", "SW4", "Q3", "PAJERO",
            "Q3 SPORTBACK", "XC90", "BRONCO"],
    "Picape": ["STRADA", "TORO", "SAVEIRO", "MONTANA", "S10", "RAMPAGE", "OROCH", "RANGER", "AMAROK", "FRONTIER",
               "HILUX", "L200"],
    "Utilitário": ["SPIN", "FIORINO", "PARTNER", "MASTER", "KANGOO"],
}
CORPO = {m: c for c, ms in CARROCERIA.items() for m in ms}
LIMITE = 0.10


def mes(v):
    """Cabeçalho de coluna → 'AAAA-MM'; None para colunas repetidas ('Jul-25*') ou vazias."""
    if isinstance(v, datetime):
        return f"{v.year}-{v.month:02d}"
    if isinstance(v, str):
        m = re.fullmatch(r"([A-Za-z]{3})-(\d{2})", v.strip())
        if m and m.group(1).lower() in MESES:
            return f"20{m.group(2)}-{MESES[m.group(1).lower()]:02d}"
    return None


def num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def cabecalho(ws, linha, c0, c1):
    """Linha de datas mais próxima acima de `linha`, entre as colunas c0 e c1: {coluna: 'AAAA-MM'}."""
    for r in range(linha - 1, 0, -1):
        cols = {c: mes(ws.cell(r, c).value) for c in range(c0, c1 + 1)}
        cols = {c: m for c, m in cols.items() if m}
        if len(cols) >= 3:
            return cols
    return {}


def serie(ws, linha, c0, c1):
    cab = cabecalho(ws, linha, c0, c1)
    out = {}
    for c, m in cab.items():
        v = num(ws.cell(linha, c).value)
        if v is not None and m not in out:
            out[m] = round(v * 100, 3)
    return out


def bloco_rotulos(ws, c_rot, c0, c1, inicio, fim):
    """Séries de cada rótulo na coluna c_rot, linhas inicio..fim."""
    out = {}
    for r in range(inicio, fim + 1):
        rot = ws.cell(r, c_rot).value
        if rot is None:
            continue
        s = serie(ws, r, c0, c1)
        if s:
            out[TIPOS.get(str(rot), str(rot))] = s
    return out


def achar(ws, texto, col=None, depois=0):
    for row in ws.iter_rows(min_row=depois + 1):
        for c in row:
            if c.value == texto and (col is None or c.column == col):
                return c.row, c.column
    return None


def modelos(ws):
    """Blocos 'By Each Car Model': título do tipo com as datas na mesma linha; modelos embaixo."""
    out = {}
    for row in ws.iter_rows():
        for c in row:
            if str(c.value) in TIPOS and c.column > 20:
                datas = {}
                for k in range(1, 8):
                    m = mes(ws.cell(c.row, c.column + k).value)
                    if m:
                        datas[c.column + k] = m
                r, vazias = c.row + 1, 0
                while vazias < 4 and r <= ws.max_row:  # a lista pode ter linha em branco no meio
                    v = ws.cell(r, c.column).value
                    if v is None or str(v).startswith("="):
                        vazias += 1
                        r += 1
                        continue
                    if str(v) in TIPOS:
                        break
                    vazias = 0
                    nome = str(v).strip().upper()
                    vals = {m: round(num(ws.cell(r, col).value) * 100, 3) for col, m in datas.items()
                            if num(ws.cell(r, col).value) is not None}
                    out[nome] = {"tipo": TIPOS[str(c.value)], "carroceria": CORPO.get(nome, "Outros"), "var": vals}
                    r += 1
    return out


def medias(mods):
    meses = sorted({m for d in mods.values() for m in d["var"]})
    out = {}
    for corpo in list(CARROCERIA) + ["Outros", "Todos"]:
        s = {}
        for m in meses:
            v = [d["var"][m] for d in mods.values() if m in d["var"] and abs(d["var"][m]) <= LIMITE * 100
                 and (corpo == "Todos" or d["carroceria"] == corpo)]
            if v:
                s[m] = {"media": round(statistics.fmean(v), 3), "n": len(v)}
        if s:
            out[corpo] = s
    return out


def juntar(velho, novo):
    """Mescla recursiva: séries {mês: valor} somam meses; o valor novo vence no mesmo mês."""
    if isinstance(velho, dict) and isinstance(novo, dict):
        out = dict(velho)
        for k, v in novo.items():
            out[k] = juntar(velho.get(k), v) if k in velho else v
        return out
    return novo


def main(xlsx, saida, base=None):
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ld, fp = wb["Cover - Listing Delta"], wb["Cover - FIPE"]
    r9 = achar(ld, "Consolidated", 2, 7)[0]
    rt = achar(ld, "By Car type", 2)[0]
    ry = achar(ld, "By Model Year", 2)[0]
    listing = {
        "consolidado": serie(ld, r9, 3, 26),
        "tipo": bloco_rotulos(ld, 2, 3, 26, rt + 2, ry - 1),
        "ano_modelo": bloco_rotulos(ld, 2, 3, 26, ry + 2, ry + 8),
    }
    mods = modelos(ld)
    rc = achar(fp, "Consolidated", 2, 8)[0]
    rft = achar(fp, "By Car type", 2)[0]
    rfy = achar(fp, "By Model Year", 2)[0]
    ra = achar(fp, "Average FIPE MoM Delta ", 2)[0]
    spread = {
        "consolidado": serie(fp, rc, 3, 26),
        "tipo": bloco_rotulos(fp, 2, 3, 26, rft + 2, rfy - 1),
        "ano_modelo": bloco_rotulos(fp, 2, 3, 26, rfy + 2, ra - 1),
    }
    rac = achar(fp, "Consolidated", 2, ra + 2)[0]
    ract = achar(fp, "By Car type", 2, ra)[0]
    fipe = {"consolidado": serie(fp, rac, 3, 26), "tipo": bloco_rotulos(fp, 2, 3, 26, ract + 2, ract + 5)}
    doc = {
        "geradoEm": datetime.now().isoformat(timespec="seconds"),
        "fonte": "Seminovos Tracker (planilha enviada pelo usuário) · preços anunciados da Localiza e FIPE",
        "arquivo": xlsx.split("/")[-1],
        "listing": listing, "modelos": mods, "carroceria": medias(mods), "spread_fipe": spread, "fipe": fipe,
        "nota": f"Variações em % no mês. Média por carroceria = média simples dos modelos, sem variações acima de "
                f"{int(LIMITE * 100)}% em módulo.",
    }
    if base:
        with open(base, encoding="utf-8") as f:
            ant = json.load(f)
        ant = ant.get("data", ant)
        doc = {**juntar({k: ant[k] for k in ("listing", "modelos", "spread_fipe", "fipe") if k in ant}, doc), **{
            k: doc[k] for k in ("geradoEm", "fonte", "arquivo", "nota")}}
        doc["carroceria"] = medias(doc["modelos"])
    with open(saida, "w", encoding="utf-8") as f:
        json.dump(doc, f, ensure_ascii=False, indent=1)
    print(f"listing {len(listing['consolidado'])} meses · {len(mods)} modelos · FIPE {len(fipe['consolidado'])} meses")


if __name__ == "__main__":
    args = sys.argv[1:]
    base = args[args.index("--base") + 1] if "--base" in args else None
    main(args[0], args[1], base)
